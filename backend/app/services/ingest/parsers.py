"""上傳文檔解析:把各種檔案格式抽成純文字,接回既有「切塊→嵌入→入庫」管線。

支援:.pdf(pypdf)、.docx(python-docx)、.txt/.md 及其他可 UTF-8 解碼的文字檔。
無法解析(格式不支援 / 加密 / 內容損毀 / 空白)一律丟 ``DocumentParseError``,
由 API 層轉成清楚的 4xx 錯誤,不讓例外冒成 500。

相依(pypdf、python-docx)採「函式內延遲匯入」:未安裝時給明確訊息,
且不影響 app 其餘部分啟動(對齊 app.services.embedding 對 fastembed 的做法)。
"""
from __future__ import annotations

import io
from pathlib import Path


class DocumentParseError(Exception):
    """文檔解析失敗(格式不支援 / 加密 / 損毀 / 內容為空)。API 層應轉成 4xx。"""


# 可直接以純文字處理的副檔名(其餘未知者亦嘗試 UTF-8 解碼)
_TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".text", ".csv", ".log", ".rst"}
# 對前端 / 使用者提示用的「明確支援」清單
SUPPORTED_DOC_EXTENSIONS = {".pdf", ".docx"} | _TEXT_EXTENSIONS


def _ext(filename: str | None) -> str:
    return Path(filename or "").suffix.lower()


def parse_pdf(raw: bytes) -> str:
    """用 pypdf 逐頁抽取文字。加密(需密碼)或掃描圖片(抽不到字)給明確錯誤。"""
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as e:  # pragma: no cover - 相依未裝時
        raise DocumentParseError("伺服器缺少 PDF 解析相依(pypdf),請聯絡管理員安裝") from e

    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as e:  # 損毀 / 非法 PDF
        raise DocumentParseError("PDF 檔無法讀取(可能損毀或非有效 PDF)") from e

    if reader.is_encrypted:
        # 嘗試空密碼解密;失敗(回 0 / 例外)代表需要密碼,無法解析。
        try:
            if reader.decrypt("") == 0:
                raise DocumentParseError("PDF 已加密,需密碼才能解析;請提供未加密版本")
        except DocumentParseError:
            raise
        except Exception as e:
            raise DocumentParseError("PDF 已加密且無法解密;請提供未加密版本") from e

    parts: list[str] = []
    try:
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - 單頁抽取失敗跳過,不拖垮整份
                continue
    except Exception as e:
        raise DocumentParseError("PDF 內容抽取失敗") from e

    text = "\n\n".join(p.strip() for p in parts if p and p.strip())
    if not text.strip():
        raise DocumentParseError("PDF 未能抽出任何文字(可能為掃描圖片,需 OCR,暫不支援)")
    return text


def parse_docx(raw: bytes) -> str:
    """用 python-docx 抽段落 + 表格文字。"""
    try:
        import docx  # python-docx 匯入名為 docx
    except ModuleNotFoundError as e:  # pragma: no cover - 相依未裝時
        raise DocumentParseError(
            "伺服器缺少 DOCX 解析相依(python-docx),請聯絡管理員安裝"
        ) from e

    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as e:  # 損毀 / 非 zip 結構(如舊版 .doc)
        raise DocumentParseError("DOCX 檔無法讀取(可能損毀或非有效 .docx)") from e

    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # 併入表格文字(逐列以 tab 分隔),常見於產品規格 / FAQ 表
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise DocumentParseError("DOCX 未能抽出任何文字(檔案可能為空)")
    return text


def parse_text(raw: bytes) -> str:
    """嘗試以 UTF-8(含 BOM)解碼純文字檔。"""
    text: str | None = None
    for enc in ("utf-8", "utf-8-sig"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            text = None
    if text is None:
        raise DocumentParseError(
            "僅支援 UTF-8 文字檔;請另存為 UTF-8,或改用 PDF / DOCX 格式"
        )
    if not text.strip():
        raise DocumentParseError("檔案內容為空")
    return text


def parse_document(filename: str | None, raw: bytes) -> str:
    """依副檔名把上傳檔內容抽成純文字。無法解析時丟 ``DocumentParseError``。

    - ``.pdf`` → pypdf;``.docx`` → python-docx;``.txt/.md`` 等 → UTF-8 解碼。
    - ``.doc``(舊二進位 Word)非 zip 結構,明確擋掉並引導改存 .docx。
    - 其餘未知副檔名一律嘗試以純文字(UTF-8)處理。
    """
    if not raw:
        raise DocumentParseError("檔案為空")
    ext = _ext(filename)
    if ext == ".pdf":
        return parse_pdf(raw)
    if ext == ".docx":
        return parse_docx(raw)
    if ext == ".doc":
        raise DocumentParseError("不支援舊版 .doc;請另存為 .docx 或 PDF 後再上傳")
    return parse_text(raw)
